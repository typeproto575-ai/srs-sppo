from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.enum.text import WD_COLOR_INDEX


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "SRS-SPPO_User_Manual.docx"
OUTPUT = ROOT / "docs" / "SRS-SPPO_User_Manual_v1.1.docx"
STYLE_MAP: dict[str, object] = {}


def paragraph_text(paragraph: Paragraph) -> str:
    return paragraph.text


def find_paragraph(document: Document, text: str) -> Paragraph:
    matches = [p for p in document.paragraphs if paragraph_text(p) == text]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph matching {text!r}; found {len(matches)}.")
    return matches[0]


def find_in_section(
    document: Document,
    heading: str,
    text_prefix: str,
) -> Paragraph:
    paragraphs = document.paragraphs
    start = next((i for i, p in enumerate(paragraphs) if p.text == heading), None)
    if start is None:
        raise ValueError(f"Section heading not found: {heading!r}")
    end = next(
        (
            i
            for i in range(start + 1, len(paragraphs))
            if paragraphs[i].style.style_id == "Heading1"
        ),
        len(paragraphs),
    )
    matches = [
        p for p in paragraphs[start + 1 : end] if p.text.startswith(text_prefix)
    ]
    if len(matches) != 1:
        raise ValueError(
            f"Expected one paragraph starting with {text_prefix!r} "
            f"in {heading!r}; found {len(matches)}."
        )
    return matches[0]


def insert_after(
    anchor: Paragraph,
    text: str,
    style_name: str,
) -> Paragraph:
    new_element = OxmlElement("w:p")
    anchor._p.addnext(new_element)
    paragraph = Paragraph(new_element, anchor._parent)
    paragraph.style = STYLE_MAP[style_name]
    if text:
        paragraph.add_run(text)
    return paragraph


def replace_text(paragraph: Paragraph, text: str) -> None:
    paragraph.clear()
    paragraph.add_run(text)


def insert_sequence(
    anchor: Paragraph,
    items: list[tuple[str, str]],
) -> Paragraph:
    for style_name, text in items:
        anchor = insert_after(anchor, text, style_name)
    return anchor


def insert_screenshot_placeholder(anchor: Paragraph, description: str) -> Paragraph:
    placeholder = insert_after(anchor, f"[Insert screenshot: {description}]", "Caption")
    for run in placeholder.runs:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return placeholder


def main() -> None:
    global STYLE_MAP

    if not SOURCE.is_file():
        raise FileNotFoundError(f"Original manual not found: {SOURCE}")
    if OUTPUT.exists():
        raise FileExistsError(
            f"Output already exists; refusing to overwrite it: {OUTPUT}"
        )

    document = Document(SOURCE)
    STYLE_MAP = {style.name: style for style in document.styles}
    change_log = [
        "Cover: edited the version label from Version 1.0 to Version 1.1.",
        "Who should use each part and Chapter 1: edited role guidance to clarify the single Secretary staff role and technical-administrator responsibilities.",
        "Chapter 6: added reset, lockout, and idle-timeout guidance; replaced the dashboard menu summary with source-verified navigation and notification details.",
        "Contents and Chapter 7: added a correction-request section and its register-check safeguards.",
        "Chapter 9: replaced the generic six-step certificate workflow with seven source-verified steps; added the confirmed regeneration control while retaining the preview and in-person release notes.",
        "Chapter 11: edited account, database, backup, restore, and person-tool guidance; added Email Delivery Log and Clean Up Old Data sections with safeguards.",
        "Chapter 12: replaced the generic login-help bullet with specific idle logout, failed-attempt, and reset-email checks.",
        "Added eight highlighted screenshot placeholders in the existing Caption style (no screenshots created):",
    ]
    screenshot_placeholders = [
        "[Insert screenshot: Chapter 6 dashboard, top bar and navigation groups]",
        "[Insert screenshot: Certificates > Print / Pickup History, release form]",
        "[Insert screenshot: Certificates > Correction Requests, request review and action buttons]",
        "[Insert screenshot: System > Accounts, staff account controls and password rule]",
        "[Insert screenshot: System > Backup, encrypted export form]",
        "[Insert screenshot: System > Database Setup, update statuses and actions]",
        "[Insert screenshot: System > Email Log, delivery status list]",
        "[Insert screenshot: System > Clean Up Old Data, preview counts and confirmation]",
    ]

    # Cover
    replace_text(
        find_paragraph(document, "Draft for review | Version 1.0"),
        "Version 1.1",
    )

    # Audience and role guidance
    replace_text(
        find_paragraph(
            document,
            "•  System Administrator: Chapter 11 covers accounts, backups, database updates, and routine system safeguards. In the current system, staff accounts use the Secretary role; a separate administrator role is not provided.",
        ),
        "•  System Administrator: Chapter 11 covers system-maintenance tasks. There is only one staff role, Secretary; in-app tasks use the same Secretary accounts, while hosting, recovery, and database restoration are handled by the designated technical administrator.",
    )
    replace_text(
        find_paragraph(
            document,
            "•  System maintainer: protects staff access, applies approved database updates, and creates protected backups. Ask the designated technical administrator for hosting or recovery actions.",
        ),
        "•  System maintainer: uses the same Secretary staff account for in-app maintenance, protects staff access, applies approved database updates, and creates protected backups. The designated technical administrator handles hosting, recovery, and database restoration.",
    )

    # Chapter 6: Sign in
    sign_in_heading = "6. Secretary sign-in and dashboard"
    step_two = find_in_section(
        document, sign_in_heading, "2.  Enter the assigned username and password"
    )
    step_three = find_in_section(
        document, sign_in_heading, "3.  If sign-in fails"
    )
    replace_text(
        step_three,
        "4.  If sign-in is blocked, 8 failed attempts for one username within 15 minutes or 20 failed attempts from one IP address within 15 minutes blocks further attempts for 15 minutes. Wait 15 minutes, then use \"Forgot your password?\" first. If that does not resolve access, ask the designated technical administrator or another active staff member.",
    )
    insert_after(
        step_two,
        "3.  If you forget your password, select \"Forgot your password?\" on the login page, enter the address in the \"Staff email\" field, select \"Send reset link,\" and use the emailed link. The link works once and expires after 30 minutes. The reset form requires at least 12 characters; it does not currently enforce uppercase, lowercase, and number requirements.",
        "List Paragraph",
    )
    insert_after(
        step_three,
        "Staff sessions end after 30 minutes of inactivity. Sign in again to continue.",
        "Note",
    )

    # Chapter 6: Dashboard navigation
    dashboard_heading = "6. Secretary sign-in and dashboard"
    dashboard_updates = {
        "•  Records:": "•  Records: All Records, Baptismal, Confirmation, First Communion, Marriage, Person Linking, Persons.",
        "•  Certificates:": "•  Certificates: Generate / Preview, Print / Pickup History, Correction Requests.",
        "•  Schedule:": "•  Schedule: Calendar, Booking Conflicts, Announcements.",
        "•  Reports:": "•  Reports: top-bar link (not a sidebar item) for monthly/yearly activity.",
        "•  System:": "•  System: Accounts, Backup, Archive / Restore, Email Log, Audit Log, Database Setup, Clean Up Old Data.",
    }
    for prefix, text in dashboard_updates.items():
        replace_text(find_in_section(document, dashboard_heading, prefix), text)
    reports = find_in_section(
        document, dashboard_heading, "•  Reports: top-bar link"
    )
    insert_after(
        reports,
        "•  Notifications: top-bar bell with \"View All\" and \"Mark all read.\"",
        "List Paragraph",
    )
    dashboard_placeholder_anchor = find_in_section(
        document,
        dashboard_heading,
        "•  Notifications: top-bar bell",
    )
    insert_screenshot_placeholder(
        dashboard_placeholder_anchor,
        "Chapter 6 dashboard, top bar and navigation groups",
    )

    # Chapter 7: Correction requests
    correction_anchor = find_paragraph(
        document,
        "8.  When ready, notify the requester according to office practice. Verify identity and record pickup/release through the certificate history workflow.",
    )
    contents_anchor = find_paragraph(
        document,
        "•  7. Reviewing certificate requests",
    )
    insert_after(
        contents_anchor,
        "•  Processing correction requests (Chapter 7)",
        "List Paragraph",
    )
    correction_end = insert_sequence(
        correction_anchor,
        [
            ("Heading 2", "Processing correction requests"),
            (
                "Normal",
                "Open Certificates > Correction Requests and open the request. Compare the requested change with the physical register before taking action.",
            ),
            (
                "List Paragraph",
                "1.  Select \"Mark as Processing\" while reviewing, or \"Reject Request\" when the request should not proceed. An internal Admin Review Note can be saved with the action.",
            ),
            (
                "List Paragraph",
                "2.  If the correction is approved, correct the record through the parish's record-correction process first. Only after the record has actually been corrected, select \"Mark as Completed.\"",
            ),
            (
                "Note",
                "A correction request does not change the parish register by itself. The staff action updates the request status and note; it does not perform the record correction.",
            ),
        ],
    )
    insert_screenshot_placeholder(
        correction_end,
        "Certificates > Correction Requests, request review and action buttons",
    )

    # Chapter 9: Certificate generation and release
    chapter_nine = "9. Generating certificates and recording pickup"
    old_steps = [
        "1.  Open Certificates > Generate / Preview and select the approved request.",
        "2.  Confirm the request is linked to the correct sacramental record and that the source entry has been verified as required.",
        "3.  Review every name, date, register reference, and printed field in the preview. Do not print if the record or certificate layout is uncertain.",
        "4.  Select the print mode approved for the parish printer and certificate stock. Use the test/preview print first when aligning a new printer or stock.",
        "5.  Print the certificate and complete the parish's required signature/seal process.",
        "6.  When the certificate is handed to the verified recipient, record the pickup/release in Print / Pickup History.",
    ]
    replacement_steps = [
        "1.  Open Certificates > Generate / Preview and select the approved request. If it is not already matched, match it to the correct digital record.",
        "2.  Compare the digital entry with the physical register book. If the matched record has not been verified, select \"Confirm Verification.\" \"Mark as Generated\" remains disabled until verification is confirmed.",
        "3.  Review the preview. For a Baptismal certificate, select \"Print on Blank Stock\" for pre-printed stock or \"Print Full Certificate\" for plain paper; use plain paper for the first alignment test. For other certificates, use \"Print Certificate\" in the print view.",
        "4.  Select \"Mark as Generated.\"",
        "5.  Open Certificates > Print / Pickup History. After the physical copy is printed, select \"Mark as Printed.\"",
        "6.  After the certificate is signed and sealed, select \"Finalize Signed and Sealed.\"",
        "7.  At pickup, verify the claimant's valid ID in person. Enter the claimant name and relationship to the subject, check \"I verified valid ID at the counter,\" add optional notes if needed, then select \"Confirm Release.\"",
    ]
    for old_text, new_text in zip(old_steps, replacement_steps):
        replace_text(find_paragraph(document, old_text), new_text)
    insert_after(
        find_paragraph(document, replacement_steps[5]),
        replacement_steps[6],
        "List Paragraph",
    )
    insert_screenshot_placeholder(
        find_paragraph(document, replacement_steps[6]),
        "Certificates > Print / Pickup History, release form",
    )
    regenerate_anchor = find_paragraph(
        document,
        "•  Ask a supervisor or technical administrator for help if generation is blocked or the print alignment is incorrect.",
    )
    insert_after(
        regenerate_anchor,
        "•  If the certificate must be generated again, select \"Regenerate\" and confirm with \"Confirm Regeneration.\" The current generation log is updated rather than silently creating another active log.",
        "List Paragraph",
    )

    # Chapter 11: Staff accounts and system maintenance
    chapter_eleven = "11. System administration and backups"
    chapter_heading = find_paragraph(document, chapter_eleven)
    insert_after(
        chapter_heading,
        "These tasks do not use a separate administrator role. In-app maintenance is done through the same Secretary staff accounts; the designated technical administrator handles hosting, recovery, and database restoration.",
        "Normal",
    )
    accounts_paragraph = find_in_section(
        document, chapter_eleven, "Open System > Accounts"
    )
    replace_text(
        accounts_paragraph,
        "Open System > Accounts to manage authorized staff accounts. Passwords set or changed here must be at least 12 characters and include an uppercase letter, a lowercase letter, and a number. An account cannot deactivate or delete itself, and the last active account cannot be deactivated or deleted. If system history references an account, deletion is blocked; deactivate that account instead.",
    )
    insert_screenshot_placeholder(
        accounts_paragraph,
        "System > Accounts, staff account controls and password rule",
    )
    replace_text(
        find_in_section(document, chapter_eleven, "3.  Apply only an update marked"),
        "3.  Check the status: \"Ready to update,\" \"Already up to date,\" or \"Needs administrator review.\" Use \"Apply update\" only when an update is ready, or \"Mark as already updated\" when that option is offered for an update already present. Stop and ask the technical administrator if review is needed.",
    )
    replace_text(
        find_in_section(document, chapter_eleven, "Open System > Backup."),
        "Open System > Backup. The password field is labeled \"Administrator password\"; enter the signed-in staff account password there. Provide a backup passphrase of at least 12 characters, then download the encrypted export and store its passphrase separately. The export excludes password hashes, password-reset tokens, submission rate-limit data, and email delivery logs.",
    )
    backup_paragraph = find_in_section(
        document, chapter_eleven, "Open System > Backup."
    )
    insert_screenshot_placeholder(
        backup_paragraph,
        "System > Backup, encrypted export form",
    )
    database_note = find_in_section(
        document,
        chapter_eleven,
        "The Public calendar event details update adds",
    )
    replace_text(
        find_in_section(
            document,
            chapter_eleven,
            "•  Test restore procedures only",
        ),
        "•  Database restoration is handled by the designated technical administrator. The command-line script scripts/decrypt-backup.php decrypts an encrypted export to a .sql file; it does not import or restore that file into the database.",
    )
    insert_screenshot_placeholder(
        database_note,
        "System > Database Setup, update statuses and actions",
    )
    backup_restore = find_in_section(
        document,
        chapter_eleven,
        "•  Database restoration is handled",
    )
    cleanup_end = insert_sequence(
        backup_restore,
        [
            ("Heading 2", "Email Delivery Log"),
            (
                "Normal",
                "Open System > Email Log to review sent, failed, and other recorded email-delivery attempts. A \"sent\" status means the configured local mail relay accepted the message; it does not guarantee delivery to the recipient's inbox.",
            ),
            ("Heading 2", "Clean Up Old Data"),
            (
                "Normal",
                "This cleanup anonymizes older finished certificate requests (released, rejected, or cancelled), sacrament applications (completed, rejected, or cancelled), and correction requests (completed or rejected), using the retention periods configured for this installation. It scrubs old audit-log details and removes older email-delivery logs, notifications, rate-limit entries, and expired or used password-reset tokens. Open submissions and permanent church-register and person records are not changed.",
            ),
            (
                "Normal",
                "Back up the database before continuing. The page first previews counts without changing data; \"Run cleanup now\" is an irreversible action. If the readiness check fails or you are unsure, stop and ask the designated technical administrator.",
            ),
        ],
    )
    email_log_paragraph = find_in_section(
        document,
        chapter_eleven,
        "Open System > Email Log",
    )
    insert_screenshot_placeholder(
        email_log_paragraph,
        "System > Email Log, delivery status list",
    )
    cleanup_paragraph = find_in_section(
        document,
        chapter_eleven,
        "Back up the database before continuing.",
    )
    insert_screenshot_placeholder(
        cleanup_paragraph,
        "System > Clean Up Old Data, preview counts and confirmation",
    )
    replace_text(
        find_in_section(document, chapter_eleven, "Use the Audit Log"),
        "Use the Audit Log to review recorded staff actions. Person Linking and Persons are under Records, not System; use these tools only after confirming the person and register entries. Database maintenance, hosting configuration, and recovery should be performed by the designated technical administrator.",
    )

    # Chapter 12: Staff troubleshooting
    chapter_twelve = "12. Troubleshooting and quick reference"
    replace_text(
        find_in_section(document, chapter_twelve, "•  Login or permission problem:"),
        "•  Logged out unexpectedly: staff sessions end after 30 minutes of inactivity; sign in again.",
    )
    staff_checks = find_in_section(
        document, chapter_twelve, "•  Logged out unexpectedly:"
    )
    insert_sequence(
        staff_checks,
        [
            (
                "List Paragraph",
                "•  Too many failed attempts: wait 15 minutes, then use \"Forgot your password?\".",
            ),
            (
                "List Paragraph",
                "•  Password reset email did not arrive: check System > Email Log, check the spam folder, and confirm the staff email address on the account in System > Accounts.",
            ),
            (
                "List Paragraph",
                "•  Login or permission problem: after trying password reset, ask the designated technical administrator or another active staff member. Do not share passwords through chat or email.",
            ),
        ],
    )

    document.save(OUTPUT)
    print(f"Updated manual saved to: {OUTPUT}")
    print("Change log:")
    for entry in change_log:
        print(f"- {entry}")
    for placeholder in screenshot_placeholders:
        print(f"  {placeholder}")


if __name__ == "__main__":
    main()
